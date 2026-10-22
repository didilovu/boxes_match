"""Проверки кода, который пользователь запускает в блокноте."""
import json
from pathlib import Path
import unittest

notebook = json.loads((Path(__file__).resolve().parents[1] / 'shlop.ipynb').read_text(encoding='utf-8'))
namespace = {}
exec(''.join(notebook['cells'][1]['source']), namespace)
pd = namespace['pd']
optimize = namespace['optimize_whole_bins']


def row(bin_id, code, group='G', sector='S', qty=1, capacity=10, address='B3-01.02.03'):
    return dict(zip(['Тара', 'Код', 'Группа', 'Сектор', 'Количество', 'Вместимость', 'Адрес'],
                    [bin_id, code, group, sector, qty, capacity, address]))


class ConsolidationTests(unittest.TestCase):
    def run_plan(self, rows, **kwargs):
        return optimize(pd.DataFrame(rows), **kwargs)

    def test_code_beats_group_and_sector_even_on_neighbor_street(self):
        result, _, _, _ = self.run_plan([
            row(1, 'a'), row(2, 'b', qty=8),
            row(3, 'a', group='other', qty=8, address='B3-02.01.01')])
        self.assertEqual(result.loc[result['Тара']==1, 'Целевая_тара'].iloc[0], 3)

    def test_group_beats_sector(self):
        result, _, _, _ = self.run_plan([row(1, 'a'), row(2, 'b', group='other', qty=8), row(3, 'c', qty=7)])
        self.assertEqual(result.loc[result['Тара']==1, 'Целевая_тара'].iloc[0], 3)

    def test_no_match_no_move(self):
        _, freed, _, _ = self.run_plan([row(1, 'a'), row(2, 'b', sector='other', qty=8)])
        self.assertFalse(freed)

    def test_floor_and_distance(self):
        for address in ['B4-01.02.03', 'B3-03.02.03']:
            _, freed, _, _ = self.run_plan([row(1, 'a'), row(2, 'a', qty=8, address=address)])
            self.assertFalse(freed)
        _, freed, _, _ = self.run_plan([row(1, 'a'), row(2, 'a', qty=8, address='B3-02.02.03')], max_street_distance=0)
        self.assertFalse(freed)

    def test_fractional_capacity_rejects_overflow(self):
        _, freed, _, _ = self.run_plan([row(1, 'a', qty=2, capacity=4), row(2, 'b', qty=8, capacity=10)])
        self.assertFalse(freed)

    def test_fractional_capacity_exact_fit(self):
        result, freed, _, groups = self.run_plan([row(1, 'a', qty=2, capacity=10), row(2, 'b', qty=8, capacity=10)])
        self.assertEqual(len(freed), 1)
        self.assertAlmostEqual(next(iter(groups.values()))['fill_fraction'], 1)
        self.assertEqual(result['Количество'].sum(), 10)

    def test_no_chains_and_target_rows_in_group(self):
        result, freed, _, _ = self.run_plan([row(1, 'a'), row(2, 'a'), row(3, 'a')])
        self.assertFalse(set(result.loc[result['Нужно_переместить'], 'Целевая_тара']) & freed)
        self.assertTrue((result['Группа_объединения'] > 0).all())

    def test_invalid_input(self):
        for bad in [row(1, 'a', capacity=0), row(1, 'a', qty=-1), row(1, 'a', address='???')]:
            with self.assertRaises(ValueError):
                self.run_plan([bad])

    def test_short_address(self):
        self.assertEqual(namespace['parse_address']('б1.02.03'), ('Б1', 2))


    def test_warehouse_sheet_has_only_movement_groups_and_full_totals(self):
        import tempfile
        from openpyxl import load_workbook
        result, _, _, groups = self.run_plan([
            row(1, 'a', qty=2), row(2, 'a', qty=7),
            row(3, 'z', sector='unrelated', qty=8)])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'report.xlsx'
            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                namespace['write_warehouse_sheet'](writer, result, groups)
            book = load_workbook(path)
            self.assertEqual(book.sheetnames[0], 'Для склада')
            sheet = book['Для склада']
            self.assertEqual(list(sheet.values), [
                ('признак объединения', 'Адрес', 'Тара', 'Сумма по количеству', 'куда'),
                (1, None, None, 9, None),
                (1, 'B3-01.02.03', 2, 7, 'да'),
                (1, 'B3-01.02.03', 1, 2, None)])
            self.assertEqual(sheet['A2'].fill.fgColor.rgb, '00BFBFBF')
            self.assertEqual(sheet['A3'].fill.fgColor.rgb, '00E2F0D9')
            self.assertEqual(sheet.freeze_panes, 'A2')
            book.close()

    def test_warehouse_sheet_without_moves_has_only_header(self):
        import tempfile
        from openpyxl import load_workbook
        result, _, _, groups = self.run_plan([row(1, 'a')])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'empty.xlsx'
            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                namespace['write_warehouse_sheet'](writer, result, groups)
            book = load_workbook(path)
            self.assertEqual(book['Для склада'].max_row, 1)
            book.close()


if __name__ == '__main__':
    unittest.main()
